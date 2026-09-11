# -*- coding: utf-8 -*-
"""자리배치도 샘플 생성 (가상 데이터)"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 팀별 색상 (배경)
TEAM_COLORS = {
    "감리단장": "FFC7CE",
    "건축팀":   "DDEBF7",
    "설비팀":   "E2EFDA",
    "전기팀":   "FCE4D6",
    "공무팀":   "E4DFEC",
    "빈자리":   "F2F2F2",
}

# 가상 좌석 데이터: (좌석번호, 이름, 팀) — 6열 x 5행
ROWS = [
    [("A1", "김감리", "감리단장"), ("A2", "", "빈자리"), ("A3", "박건축", "건축팀"), ("A4", "이건축", "건축팀"), ("A5", "최건축", "건축팀"), ("A6", "정건축", "건축팀")],
    [("B1", "오설비", "설비팀"), ("B2", "강설비", "설비팀"), ("B3", "조설비", "설비팀"), ("B4", "윤설비", "설비팀"), ("B5", "장설비", "설비팀"), ("B6", "임설비", "설비팀")],
    [("C1", "한전기", "전기팀"), ("C2", "서전기", "전기팀"), ("C3", "권전기", "전기팀"), ("C4", "황전기", "전기팀"), ("C5", "", "빈자리"), ("C6", "배전기", "전기팀")],
    [("D1", "심공무", "공무팀"), ("D2", "노공무", "공무팀"), ("D3", "하공무", "공무팀"), ("D4", "곽공무", "공무팀"), ("D5", "", "빈자리"), ("D6", "문공무", "공무팀")],
    [("E1", "신입사", "건축팀"), ("E2", "안건축", "건축팀"), ("E3", "", "빈자리"), ("E4", "유설비", "설비팀"), ("E5", "남설비", "설비팀"), ("E6", "", "빈자리")],
]

wb = Workbook()
ws = wb.active
ws.title = "자리배치도"

thin = Side(style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 제목
ws.merge_cells("A1:F1")
c = ws["A1"]
c.value = "P4 현장 사무실 자리배치도  (샘플 — 가상 인원)"
c.font = Font(size=14, bold=True)
c.alignment = center
ws.row_dimensions[1].height = 30

# 범례
ws.merge_cells("A2:F2")
legend = " | ".join(f"{t}" for t in TEAM_COLORS if t != "빈자리")
c = ws["A2"]
c.value = "범례: " + legend
c.font = Font(size=9, color="666666")
c.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[2].height = 18

# 좌석 그리드: 4행부터, 각 좌석 2셀 높이 (이름/팀), 좌석 블록 사이에 복도 1행
start_row = 4
for r_idx, row in enumerate(ROWS):
    r1 = start_row + r_idx * 3
    r2 = r1 + 1
    for col_idx, (seat, name, team) in enumerate(row):
        col = col_idx + 1
        fill = PatternFill(start_color=TEAM_COLORS[team], end_color=TEAM_COLORS[team], fill_type="solid")
        for rr in (r1, r2):
            cell = ws.cell(row=rr, column=col)
            cell.border = border
            cell.fill = fill
            cell.alignment = center
        ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)
        ws.cell(row=r1, column=col).value = f"{seat}\n{name if name else '(빈자리)'}\n{team if name else ''}"
        ws.cell(row=r1, column=col).font = Font(size=9, bold=bool(name))
    ws.row_dimensions[r1].height = 30
    ws.row_dimensions[r2].height = 16

# 복도 표시 (좌석 블록 사이)
for i in range(len(ROWS) - 1):
    rr = start_row + i * 3 + 2
    cc = ws.cell(row=rr, column=1)
    cc.value = "┄ 복 도 ┄"
    cc.font = Font(size=8, color="999999")
    cc.alignment = center
    cc.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
    ws.row_dimensions[rr].height = 14

for col_idx in range(1, 7):
    ws.column_dimensions[get_column_letter(col_idx)].width = 14

OUT = r"C:\Users\orisi\.openclaw\workspace\자리배치도-샘플.xlsx"
wb.save(OUT)
print("saved:", OUT)
