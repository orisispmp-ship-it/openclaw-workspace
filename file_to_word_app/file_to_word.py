"""
파일을 워드 문서로 변환하는 앱 (PDF, JPG, PNG 지원)
"""

import os
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import threading

def convert_pdf_to_word(pdf_path, output_path):
    """PDF 파일을 Word 문서로 변환"""
    from pdf2docx import Converter
    cv = Converter(pdf_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()
    return True

def convert_image_to_word(image_path, output_path):
    """이미지 파일(JPG, PNG)을 OCR로 Word 문서로 변환"""
    try:
        from PIL import Image
        import pytesseract
        from docx import Document
        
        # Tesseract 경로 자동 탐색
        common_paths = [
            r'C:\Program Files\Tesseract-OCR\tesseract.exe',
            r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
        ]
        for p in common_paths:
            if os.path.exists(p):
                pytesseract.pytesseract.tesseract_cmd = p
                break
        
        img = Image.open(image_path)
        text = pytesseract.image_to_string(img, lang='kor+eng')
        
        doc = Document()
        doc.add_heading(os.path.basename(image_path), level=1)
        doc.add_paragraph(text)
        doc.save(output_path)
        return True
    except Exception as e:
        raise Exception(f"OCR 오류 (Tesseract 미설치 가능): {e}\nhttps://github.com/UB-Mannheim/tesseract/wiki 에서 설치하세요.")

class FileToWordApp:
    def __init__(self, root):
        self.root = root
        self.root.title("파일 → 워드 변환기")
        self.root.geometry("600x400")
        self.root.resizable(True, True)
        
        self.files = []
        
        # 스타일
        self.root.configure(bg='#f0f0f0')
        
        # 헤더
        header = tk.Label(root, text="📄 파일 → 워드 문서 변환기", font=('맑은 고딕', 16, 'bold'), bg='#2c3e50', fg='white', pady=10)
        header.pack(fill=tk.X)
        
        # 설명
        desc = tk.Label(root, text="PDF, JPG, PNG 파일을 Word 문서(.docx)로 변환합니다.", 
                       font=('맑은 고딕', 9), bg='#f0f0f0', fg='#555', pady=5)
        desc.pack()
        
        # 버튼 프레임
        btn_frame = tk.Frame(root, bg='#f0f0f0')
        btn_frame.pack(pady=10)
        
        self.add_btn = tk.Button(btn_frame, text="📂 파일 선택", command=self.add_files,
                                font=('맑은 고딕', 10), bg='#3498db', fg='white', padx=15, pady=5, cursor='hand2')
        self.add_btn.pack(side=tk.LEFT, padx=5)
        
        self.clear_btn = tk.Button(btn_frame, text="🗑️ 목록 지우기", command=self.clear_list,
                                  font=('맑은 고딕', 10), bg='#e74c3c', fg='white', padx=15, pady=5, cursor='hand2')
        self.clear_btn.pack(side=tk.LEFT, padx=5)
        
        # 파일 목록
        list_frame = tk.Frame(root, bg='white', bd=1, relief=tk.SUNKEN)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        scrollbar = tk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.listbox = tk.Listbox(list_frame, yscrollcommand=scrollbar.set, font=('맑은 고딕', 9),
                                 bg='white', selectbackground='#3498db')
        self.listbox.pack(fill=tk.BOTH, expand=True)
        scrollbar.config(command=self.listbox.yview)
        
        # 진행 상태
        self.progress = ttk.Progressbar(root, mode='indeterminate')
        self.progress.pack(fill=tk.X, padx=10, pady=5)
        
        self.status_label = tk.Label(root, text="", font=('맑은 고딕', 9), bg='#f0f0f0', fg='#555')
        self.status_label.pack()
        
        # 변환 버튼
        self.convert_btn = tk.Button(root, text="🔄 변환 시작", command=self.start_conversion,
                                    font=('맑은 고딕', 12, 'bold'), bg='#27ae60', fg='white', 
                                    padx=20, pady=8, cursor='hand2')
        self.convert_btn.pack(pady=10)
        
    def add_files(self):
        files = filedialog.askopenfilenames(
            title="변환할 파일 선택",
            filetypes=[("지원 파일", "*.pdf *.jpg *.jpeg *.png"), ("PDF", "*.pdf"), ("이미지", "*.jpg *.jpeg *.png"), ("모든 파일", "*.*")]
        )
        for f in files:
            if f not in self.files:
                self.files.append(f)
                name = os.path.basename(f)
                size = os.path.getsize(f)
                size_str = f"{size/1024:.1f} KB" if size < 1024*1024 else f"{size/1024/1024:.1f} MB"
                ext = os.path.splitext(f)[1].upper()
                self.listbox.insert(tk.END, f"[{ext}] {name} ({size_str})")
        
        self.update_status(f"선택된 파일: {len(self.files)}개")
    
    def clear_list(self):
        self.files.clear()
        self.listbox.delete(0, tk.END)
        self.update_status("")
    
    def update_status(self, msg):
        self.status_label.config(text=msg)
        self.root.update_idletasks()
    
    def start_conversion(self):
        if not self.files:
            messagebox.showwarning("알림", "변환할 파일을 먼저 선택하세요.")
            return
        
        thread = threading.Thread(target=self.convert_files, daemon=True)
        thread.start()
    
    def convert_files(self):
        self.convert_btn.config(state=tk.DISABLED)
        self.add_btn.config(state=tk.DISABLED)
        self.progress.start()
        
        output_dir = os.path.join(os.path.expanduser("~"), "Desktop", "변환된_워드문서")
        os.makedirs(output_dir, exist_ok=True)
        
        success = 0
        fail = 0
        
        for i, file_path in enumerate(self.files):
            try:
                name = os.path.basename(file_path)
                base, ext = os.path.splitext(name)
                ext = ext.lower()
                output_path = os.path.join(output_dir, f"{base}.docx")
                
                self.update_status(f"⏳ 변환 중 ({i+1}/{len(self.files)}): {name}")
                
                if ext == '.pdf':
                    convert_pdf_to_word(file_path, output_path)
                elif ext in ('.jpg', '.jpeg', '.png'):
                    convert_image_to_word(file_path, output_path)
                
                success += 1
                
            except Exception as e:
                fail += 1
                print(f"오류: {name} - {e}")
        
        self.progress.stop()
        self.convert_btn.config(state=tk.NORMAL)
        self.add_btn.config(state=tk.NORMAL)
        
        msg = f"✅ 변환 완료! 성공: {success}개, 실패: {fail}개"
        if success > 0:
            msg += f"\n📁 저장 위치: {output_dir}"
        self.update_status(msg)
        messagebox.showinfo("변환 완료", msg)
        
        # 저장 위치 열기
        if success > 0:
            os.startfile(output_dir)

if __name__ == "__main__":
    root = tk.Tk()
    app = FileToWordApp(root)
    root.mainloop()
